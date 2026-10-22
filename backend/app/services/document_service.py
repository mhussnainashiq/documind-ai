from io import BytesIO
from pathlib import Path
import re
import fitz
from docx import Document as DocxDocument
from fastapi import HTTPException

ALLOWED_EXTENSIONS = {".pdf", ".docx", ".txt"}

def extract_text(filename: str, raw_bytes: bytes) -> str:
    extension = Path(filename).suffix.lower()
    if extension not in ALLOWED_EXTENSIONS:
        raise HTTPException(status_code=415, detail="Unsupported file type. Upload a PDF, DOCX, or TXT file.")
    try:
        if extension == ".pdf":
            with fitz.open(stream=raw_bytes, filetype="pdf") as pdf:
                return "\n\n".join(page.get_text("text") for page in pdf).strip()
        if extension == ".docx":
            document = DocxDocument(BytesIO(raw_bytes))
            parts = [p.text for p in document.paragraphs if p.text.strip()]
            for table in document.tables:
                for row in table.rows:
                    parts.append(" | ".join(cell.text.strip() for cell in row.cells))
            return "\n".join(parts).strip()
        return raw_bytes.decode("utf-8-sig").strip()
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(status_code=400, detail="The file could not be read. It may be damaged, password-protected, or not a valid file.") from exc

def count_words(text: str) -> int:
    return len(text.split())

def split_into_chunks(text: str, chunk_size: int = 900, overlap: int = 150) -> list[str]:
    """Split text into overlapping word-based chunks for retrieval."""
    words = text.split()
    if not words:
        return []
    chunk_size = max(100, chunk_size)
    overlap = max(0, min(overlap, chunk_size // 2))
    chunks = []
    step = chunk_size - overlap
    for start in range(0, len(words), step):
        chunk = " ".join(words[start:start + chunk_size]).strip()
        if chunk:
            chunks.append(chunk)
        if start + chunk_size >= len(words):
            break
    return chunks
